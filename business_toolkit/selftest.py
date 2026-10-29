import os,hashlib,http.server,importlib.metadata,json,pathlib,shutil,subprocess,sys,tempfile,threading,time
from toolkit import validate_prices,render_docx
from docx import Document
from PIL import Image,ImageDraw,ImageFont
from fpdf import FPDF
from pdfminer.high_level import extract_text
ROOT=pathlib.Path(__file__).resolve().parent
checks=[]
def check(name,ok):
 checks.append({'test':name,'passed':bool(ok)})
 if not ok:raise AssertionError(name)
def run(cmd):return subprocess.run(cmd,capture_output=True,text=True,timeout=120,env={k:v for k,v in os.environ.items() if k.lower() not in {"http_proxy","https_proxy","all_proxy"}})
with tempfile.TemporaryDirectory() as folder:
 p=pathlib.Path(folder)
 header='supplier_id,grade,price_usd_mt,basis,source_url,checked_date\n'
 (p/'good.csv').write_text(header+'SYNTHETIC-1,SG-5,900,FCA,https://example.com,2026-10-29\n')
 (p/'bad.csv').write_text(header+'SYNTHETIC-1,UNKNOWN,-4,FCA,https://example.com,2026-10-29\n')
 check('Frictionless accepts valid sample',validate_prices(p/'good.csv')['valid'])
 check('Frictionless rejects grade and negative price',not validate_prices(p/'bad.csv')['valid'])
 doc=Document();doc.sections[0].header.paragraphs[0].text='SYNTHETIC HEADER';doc.add_paragraph('Item: {{ item }}');doc.sections[0].footer.paragraphs[0].text='SYNTHETIC FOOTER';doc.save(p/'master.docx')
 h=hashlib.sha256((p/'master.docx').read_bytes()).hexdigest();(p/'registry.json').write_text(json.dumps({h:{'status':'APPROVED','scope':'synthetic test only'}}))
 r=render_docx(p/'master.docx',{'item':'AC & equipment'},p/'registry.json',p/'draft.docx')
 d=Document(p/'draft.docx');check('DOCX render preserves header/footer and escapes text',d.paragraphs[0].text=='Item: AC & equipment' and d.sections[0].header.paragraphs[0].text=='SYNTHETIC HEADER' and d.sections[0].footer.paragraphs[0].text=='SYNTHETIC FOOTER')
 check('DOCX source unchanged and release denied',hashlib.sha256((p/'master.docx').read_bytes()).hexdigest()==h and r['release_authorized'] is False)
 (p/'registry.json').write_text('{}')
 try:render_docx(p/'master.docx',{'item':'X'},p/'registry.json',p/'bad.docx');rejected=False
 except ValueError:rejected=True
 check('Unregistered DOCX master rejected',rejected)
 im=Image.new('RGB',(1800,1100),'white');dr=ImageDraw.Draw(im);font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',65);dr.text((100,250),'SYNTHETIC INVOICE 4821',font=font,fill='black');dr.text((100,400),'TOTAL USD 900',font=font,fill='black');im.save(p/'scan.png')
 pdf=FPDF();pdf.add_page();pdf.image(str(p/'scan.png'),x=0,y=0,w=210);pdf.output(str(p/'scan.pdf'))
 r=run([sys.executable,'-m','ocrmypdf','--output-type','pdf','--optimize','0','-l','eng',str(p/'scan.pdf'),str(p/'searchable.pdf')]);check('OCRmyPDF executes actual image-only scan',r.returncode==0)
 check('OCR output has searchable invoice text','4821' in extract_text(p/'searchable.pdf'))
 binary=str(ROOT/'bin/pdfcpu');check('pdfcpu validates OCR PDF',run([binary,'validate',str(p/'searchable.pdf')]).returncode==0)
 check('pdfcpu merges PDFs',run([binary,'merge',str(p/'merged.pdf'),str(p/'searchable.pdf'),str(p/'searchable.pdf')]).returncode==0)
 (p/'broken.pdf').write_bytes(b'NOT A PDF');check('pdfcpu rejects malformed PDF',run([binary,'validate',str(p/'broken.pdf')]).returncode!=0)
 class Handler(http.server.BaseHTTPRequestHandler):
  value='SYNTHETIC PRICE 900'
  def do_GET(self):
   b=self.value.encode();self.send_response(200);self.send_header('Content-Type','text/plain');self.end_headers();self.wfile.write(b)
  def log_message(self,*args):pass
 server=http.server.ThreadingHTTPServer(('127.0.0.1',0),Handler);threading.Thread(target=server.serve_forever,daemon=True).start()
 (p/'urls.yaml').write_text('name: Synthetic local price\nurl: http://127.0.0.1:'+str(server.server_port)+'/\n')
 (p/'config.yaml').write_text('report:\n  stdout:\n    enabled: true\ndisplay:\n  new: true\n  unchanged: false\n  error: true\n')
 args=[str(pathlib.Path(sys.executable).parent/'urlwatch'),'--urls',str(p/'urls.yaml'),'--config',str(p/'config.yaml'),'--cache',str(p/'cache.db')]
 a=run(args);Handler.value='SYNTHETIC PRICE 875';b=run(args);server.shutdown()
 check('urlwatch initial snapshot and changed-page diff',a.returncode==0 and b.returncode==0 and '875' in b.stdout and '900' in b.stdout)
check('Datasette Lite static app and worker installed',all((ROOT/'vendor/datasette-lite'/n).is_file() for n in ['index.html','webworker.js','app.css','LICENSE']))
report={'date':'2026-10-29','scope':'Synthetic local acceptance tests only','checks':checks,'passed':sum(x['passed'] for x in checks),'versions':{n:importlib.metadata.version(n) for n in ['urlwatch','docxtpl','frictionless','ocrmypdf']},'automatic_monitoring_enabled':False,'external_release_authorized':False}
(ROOT/'test-report.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
